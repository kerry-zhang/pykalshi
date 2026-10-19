"""
Market volume dashboard.

Pulls settled and open market data, maps to categories via series lookup,
and exports an Excel dashboard showing volume by category over time.
Also includes a trade P&L sheet built from trades.csv.
"""

import sys
import os
import csv
from collections import defaultdict
from datetime import datetime, timezone

import httpx
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference

from urllib.parse import urlencode

from pykalshi import KalshiClient, MarketStatus
from pykalshi.models import MarketModel as _MarketModel

TRADES_CSV   = os.path.join(os.path.dirname(__file__), "trades.csv")
MARKET_CACHE = os.path.join(os.path.dirname(__file__), "market_status_cache.json")

EXCEL_OUT = os.path.join(os.path.dirname(__file__), "market_dashboard.xlsx")
SETTLED_PAGES = 10        # pages of 200 = up to 2,000 settled markets
OPEN_PAGES    = 20        # pages of 200 = up to 4,000 open/closed markets for 24h snapshot
MIN_VOLUME    = 1.0       # drop truly zero-volume markets


# ---------------------------------------------------------------------------
# Data collection
# ---------------------------------------------------------------------------

_SERIES_SEED = {
    # Tennis
    "KXATPMATCH": "Sports", "KXATPCHALLENGERMATCH": "Sports",
    "KXITFMATCH": "Sports", "KXITFWMATCH": "Sports", "KXWTAMATCH": "Sports",
    # US major leagues
    "KXNBA": "Sports", "KXNFL": "Sports", "KXMLB": "Sports", "KXMLBGAME": "Sports",
    "KXNHL": "Sports", "KXNHLOVERTIME": "Sports", "KXMMA": "Sports",
    "KXNASCAR": "Sports", "KXPGA": "Sports", "KXTENNIS": "Sports",
    # Soccer / international
    "KXMLS": "Sports", "KXJLEAGUE": "Sports", "KXJLEAGUEGAME": "Sports",
    "KXBUNDESLIGA": "Sports", "KXLALIGA": "Sports", "KXPREMIERLEAGUE": "Sports",
    "KXLIGUE1": "Sports", "KXSERIEA": "Sports", "KXEREDIVIS": "Sports",
    "KXCHAMPIONSLEAGUE": "Sports", "KXEUROPALEAGUE": "Sports",
    "KXAFCON": "Sports", "KXWORLDCUP": "Sports", "KXCOPALIBERTADORES": "Sports",
    "KXCRICKET": "Sports", "KXCRICKETTESTMATCH": "Sports",
    "KXBOXING": "Sports", "KXRUGBY": "Sports", "KXGOLF": "Sports",
    # Esports
    "KXCS2": "Esports", "KXCS2GAME": "Esports",
    "KXLOL": "Esports", "KXDOTA2": "Esports", "KXVALORANT": "Esports",
    "KXOVERWATCH": "Esports", "KXROCKETLEAGUE": "Esports",
    # Crypto
    "KXBTC": "Crypto", "KXBTC15M": "Crypto",
    "KXETH": "Crypto", "KXXRPD": "Crypto",
    "KXSOLD": "Crypto", "KXAVAXD": "Crypto", "KXSUID": "Crypto",
    "KXHYPE15M": "Crypto",
    # Economics / macro
    "KXCPI": "Economics", "KXFOMC": "Economics", "KXPCE": "Economics",
    "KXJOBS": "Economics", "KXGDP": "Economics", "KXUNRATE": "Economics",
    "KXWAGESAT": "Economics", "KXGOVIANOMR": "Economics",
    # Financials
    "KXSPX": "Financials", "KXNASDAQ": "Financials", "KXDOW": "Financials",
    # Politics / elections
    "KXPOTUS": "Politics", "KXSENATE": "Politics", "KXHOUSE": "Politics",
    # Mentions (word/phrase markets — all grouped together)
    "KXTRUMPSAY": "Mentions", "KXTRUMPMENTION": "Mentions",
    "KXTRUMPMENTIONB": "Mentions", "KXVANCEMENTION": "Mentions",
    "KXHEARINGMENTION": "Mentions", "KXMRBEASTMENTION": "Mentions",
    "KXEARNINGSMENTION": "Mentions", "KXEARNINGSMENTIONTGT": "Mentions",
    # Entertainment / culture
    "KXOSCARS": "Entertainment", "KXEMMYS": "Entertainment",
    "KXGRAMMYS": "Entertainment", "KXDUBLINBY": "Entertainment",
    # IPO
    "KXIPOANTHROPIC": "IPO", "KXIPOOPENAI": "IPO",
}

# Keyword fallbacks for series not in the seed — checked in order
_KEYWORD_RULES: list[tuple[list[str], str]] = [
    (["ATP", "ITF", "WTA", "MATCH", "GAME", "NBA", "NFL", "MLB", "NHL",
      "SOCCER", "FOOTBALL", "CRICKET", "RUGBY", "MLS", "LEAGUE", "CUP",
      "CHAMPION", "TENNIS", "BOXING", "MMA", "NASCAR", "PGA", "GOLF"],  "Sports"),
    (["CS2", "LOL", "DOTA", "VALORANT", "OVERWATCH", "ESPORT", "ROCKET"], "Esports"),
    (["BTC", "ETH", "SOL", "XRP", "AVAX", "CRYPTO", "COIN", "HYPE"],     "Crypto"),
    (["CPI", "FOMC", "PCE", "GDP", "UNRATE", "JOBS", "WAGE", "FED"],      "Economics"),
    (["SPX", "NASDAQ", "DOW", "STOCK", "EQUITY"],                          "Financials"),
    (["MENTION", "SAY", "TRUMP", "VANCE", "HEARING", "MRBEAST", "EARNINGS"], "Mentions"),
    (["POTUS", "SENATE", "HOUSE", "ELECTION", "VOTE"],                    "Politics"),
]


def _guess_category(series: str) -> str:
    s = series.upper()
    for keywords, category in _KEYWORD_RULES:
        if any(k in s for k in keywords):
            return category
    return "Other"


def build_series_category_map() -> dict:
    """Return a map from series ticker → category using seed + keyword heuristics (no API calls)."""
    class HeuristicMap(dict):
        def __missing__(self, key):
            val = _guess_category(key)
            self[key] = val
            return val

    m = HeuristicMap(_SERIES_SEED)
    print(f"Series map initialised ({len(_SERIES_SEED)} seeded, unknown series use keyword heuristics).")
    return m


def infer_series(ticker: str) -> str:
    """Extract series prefix from a market ticker (e.g. KXBTC from KXBTC-26MAY15-T88799)."""
    return ticker.split("-")[0]


def pull_settled(client, series_map: dict, pages: int = SETTLED_PAGES) -> list[dict]:
    """Pull settled non-MVE markets via raw HTTP so the cursor advances correctly."""
    print(f"Fetching settled markets (up to {pages} pages x 200, MVE excluded)...")
    rows = []
    cursor = None

    for page in range(pages):
        params: dict = {"status": "settled", "limit": 200, "mve_filter": "exclude"}
        if cursor:
            params["cursor"] = cursor
        qs = urlencode({k: v for k, v in params.items() if v is not None})
        response = client.get(f"/markets?{qs}")

        markets_data = response.get("markets") or []
        cursor = response.get("cursor") or ""

        for m_data in markets_data:
            try:
                m = _MarketModel.model_validate(m_data)
            except Exception:
                continue
            vol = float(m.volume_fp or 0)
            if vol < MIN_VOLUME:
                continue

            series = m.series_ticker or infer_series(m.ticker)
            category = series_map.get(series, _guess_category(series))
            close_dt = _parse_dt(m.close_time)

            rows.append({
                "ticker":      m.ticker,
                "series":      series,
                "category":    category,
                "volume":      vol,
                "result":      (m.result or "").lower(),
                "close_month": close_dt.strftime("%Y-%m") if close_dt else "Unknown",
                "close_date":  close_dt.date() if close_dt else None,
            })

        print(f"  Page {page + 1}/{pages}: {len(markets_data)} markets, {len(rows)} kept so far")

        if not cursor or len(markets_data) < 200:
            break

    print(f"  Total settled markets with volume: {len(rows)}")
    return rows


def pull_open_24h(client, series_map: dict, pages: int = OPEN_PAGES) -> list[dict]:
    """Pull open and recently closed markets, return 24h volume by category."""
    rows = []
    for status, label in [("open", "open"), ("closed", "closed")]:
        print(f"Fetching {label} markets for 24h volume snapshot (up to {pages} pages)...")
        cursor = None
        count = 0
        for page in range(pages):
            params: dict = {"status": status, "limit": 200}
            if cursor:
                params["cursor"] = cursor
            qs = urlencode({k: v for k, v in params.items() if v is not None})
            try:
                response = client.get(f"/markets?{qs}")
            except Exception as e:
                print(f"  Warning: could not fetch {label} markets: {e}")
                break
            markets_data = response.get("markets") or []
            cursor = response.get("cursor") or ""
            for m_data in markets_data:
                try:
                    m = _MarketModel.model_validate(m_data)
                    if m.ticker.startswith("KXMVE"):
                        continue
                    vol_24h = float(m.volume_24h_fp or 0)
                    if vol_24h < MIN_VOLUME:
                        continue
                    series = m.series_ticker or infer_series(m.ticker)
                    category = series_map[series]
                    rows.append({
                        "ticker":     m.ticker,
                        "series":     series,
                        "category":   category,
                        "volume_24h": vol_24h,
                        "total_vol":  float(m.volume_fp or 0),
                    })
                    count += 1
                except Exception:
                    continue
            if not cursor or len(markets_data) < 200:
                break
        print(f"  {count} {label} markets with 24h volume > 0 (excl. MVE)")
    return rows


def _parse_dt(s: str | None) -> datetime | None:
    if not s:
        return None
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00"))
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------------

def category_totals(settled_rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(settled_rows)
    summary = (
        df.groupby("category")
          .agg(markets=("ticker", "count"), total_volume=("volume", "sum"))
          .reset_index()
          .sort_values("total_volume", ascending=False)
          .reset_index(drop=True)
    )
    summary.columns = ["Category", "Markets", "Total Volume"]
    summary["Total Volume"] = summary["Total Volume"].round(0).astype(int)
    return summary


def volume_over_time(settled_rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(settled_rows)
    pivot = (
        df.groupby(["close_month", "category"])["volume"]
          .sum()
          .reset_index()
          .pivot(index="close_month", columns="category", values="volume")
          .fillna(0)
          .astype(int)
          .reset_index()
          .rename(columns={"close_month": "Month"})
    )
    pivot = pivot.sort_values("Month").reset_index(drop=True)
    return pivot


def active_now(open_rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(open_rows)
    summary = (
        df.groupby("category")
          .agg(
              open_markets=("ticker", "count"),
              volume_24h=("volume_24h", "sum"),
              total_volume=("total_vol", "sum"),
          )
          .reset_index()
          .sort_values("volume_24h", ascending=False)
          .reset_index(drop=True)
    )
    summary.columns = ["Category", "Open Markets", "24h Volume", "Total Volume"]
    for col in ["24h Volume", "Total Volume"]:
        summary[col] = summary[col].round(0).astype(int)
    return summary


# ---------------------------------------------------------------------------
# Trade P&L
# ---------------------------------------------------------------------------

def _infer_category(ticker: str) -> str:
    """Infer category from ticker when trades.csv has no category set."""
    return _guess_category(infer_series(ticker))


def _load_market_cache() -> dict:
    if os.path.exists(MARKET_CACHE):
        import json
        with open(MARKET_CACHE) as f:
            return json.load(f)
    return {}


def _save_market_cache(cache: dict):
    import json
    with open(MARKET_CACHE, "w") as f:
        json.dump(cache, f)


ELECTIONS_EVENTS_API = "https://api.elections.kalshi.com/v1/events/"


def _lookup_elections_api(ticker: str) -> tuple[str, str] | None:
    """Fallback status/result lookup for tickers that have rolled off the
    trading API's ~2-month retention window (get_market() 404s on these)."""
    event_ticker = ticker.rsplit("-", 1)[0]
    try:
        r = httpx.get(ELECTIONS_EVENTS_API + event_ticker,
                       params={"exclude_nested_markets": "false"}, timeout=10)
        if r.status_code != 200:
            return None
        for m in r.json().get("event", {}).get("markets", []):
            if m.get("ticker_name") == ticker:
                return m.get("status", ""), (m.get("result") or "").lower()
    except Exception:
        pass
    return None


def build_trade_pnl(client) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (settled_df, open_df) from trades.csv with P&L computed."""
    if not os.path.exists(TRADES_CSV):
        return pd.DataFrame(), pd.DataFrame()

    fills = list(csv.DictReader(open(TRADES_CSV, newline="")))
    if not fills:
        return pd.DataFrame(), pd.DataFrame()

    # Group fills by ticker
    by_ticker = defaultdict(list)
    for r in fills:
        by_ticker[r["ticker"]].append(r)

    # Fetch market status — skip finalized tickers already in cache
    cache = _load_market_cache()
    results = {}
    statuses = {}
    uncached = [t for t in by_ticker if "FINALIZED" not in cache.get(t, {}).get("status", "")]
    cached   = [t for t in by_ticker if t not in uncached]
    for t in cached:
        statuses[t] = cache[t]["status"]
        results[t]  = cache[t]["result"]
    print(f"Fetching market status for {len(uncached)} tickers ({len(cached)} cached)...")
    for i, ticker in enumerate(uncached, 1):
        try:
            m = client.get_market(ticker)
            statuses[ticker] = str(m.status)
            results[ticker]  = (getattr(m, "result", None) or "").lower()
        except Exception:
            fallback = _lookup_elections_api(ticker)
            if fallback:
                statuses[ticker], results[ticker] = fallback
            else:
                statuses[ticker] = "unknown"
                results[ticker]  = ""
        if "FINALIZED" in statuses[ticker].upper():
            cache[ticker] = {"status": statuses[ticker], "result": results[ticker]}
        if i % 50 == 0:
            print(f"  {i}/{len(uncached)} fetched...")
    _save_market_cache(cache)

    def compute_position(ticker_fills):
        net_yes = 0.0
        cash_out = 0.0
        cash_in = 0.0
        fees = 0.0
        total_contracts = 0.0
        for r in ticker_fills:
            c = float(r["contracts"])
            p = float(r["yes_price"])
            fee = float(r["fee"]) if r["fee"] else 0.0
            side = r["side"]
            action = r["action"]
            fees += fee
            if side == "yes":
                if action == "buy":
                    net_yes += c; cash_out += c * p + fee; total_contracts += c
                else:
                    net_yes -= c; cash_in += c * p - fee
            else:
                if action == "buy":
                    net_yes -= c; cash_out += c * (1 - p) + fee; total_contracts += c
                else:
                    # "sell no" = Kalshi records exiting a YES position from the NO side
                    net_yes -= c; cash_in += c * p - fee
        avg_price = (cash_out - cash_in) / total_contracts if total_contracts else 0
        return net_yes, cash_out, cash_in, fees, avg_price

    settled_rows = []
    open_rows = []

    for ticker, ticker_fills in sorted(by_ticker.items()):
        net_yes, cash_out, cash_in, fees, avg_price = compute_position(ticker_fills)
        result = results.get(ticker, "")
        status = statuses.get(ticker, "")
        category = ticker_fills[0].get("category") or _infer_category(ticker)
        trade_type = ticker_fills[0].get("type", "")
        is_finalized = "FINALIZED" in status.upper()

        if is_finalized and result:
            if result == "yes":
                payout = max(net_yes, 0) * 1.0
                if net_yes < 0:
                    payout = (-net_yes) * 0.0  # held NO, lost
            else:
                payout = max(-net_yes, 0) * 1.0 if net_yes < 0 else 0.0
            pnl = payout + cash_in - cash_out
            settled_rows.append({
                "Ticker":    ticker,
                "Category":  category,
                "Type":      trade_type,
                "Result":    result.upper(),
                "Net YES":   int(net_yes),
                "Avg Price": round(avg_price, 3),
                "Cost":      round(cash_out - cash_in, 2),
                "Payout":    round(payout, 2),
                "P&L":       round(pnl, 2),
                "Fees":      round(fees, 2),
            })
        else:
            cost_basis = cash_out - cash_in
            open_rows.append({
                "Ticker":      ticker,
                "Category":    category,
                "Type":        trade_type,
                "Net YES":     int(net_yes),
                "Avg Price":   round(avg_price, 3),
                "Cost Basis":  round(cost_basis, 2),
                "Fills":       len(ticker_fills),
            })

    settled_df = pd.DataFrame(settled_rows) if settled_rows else pd.DataFrame(
        columns=["Ticker","Category","Type","Result","Net YES","Avg Price","Cost","Payout","P&L","Fees"]
    )
    open_df = pd.DataFrame(open_rows) if open_rows else pd.DataFrame(
        columns=["Ticker","Category","Type","Net YES","Avg Price","Cost Basis","Fills"]
    )

    if not settled_df.empty:
        settled_df = settled_df.sort_values("P&L", ascending=False).reset_index(drop=True)
    if not open_df.empty:
        open_df = open_df.sort_values("Cost Basis", ascending=False).reset_index(drop=True)

    return settled_df, open_df


def build_trade_analysis(settled_df: pd.DataFrame) -> dict:
    """Compute accuracy / EV analysis tables from the settled P&L DataFrame."""
    if settled_df.empty:
        return {}

    df = settled_df.copy()
    df["Direction"] = np.where(df["Net YES"] > 0, "YES",
                      np.where(df["Net YES"] < 0, "NO", "Closed"))
    df["Won"] = (
        ((df["Direction"] == "YES") & (df["Result"] == "YES")) |
        ((df["Direction"] == "NO")  & (df["Result"] == "NO"))  |
        ((df["Direction"] == "Closed") & (df["P&L"] > 0))
    )
    df["Edge"] = df["Won"].astype(float) - df["Avg Price"]

    held   = df[df["Direction"] != "Closed"]
    closed = df[df["Direction"] == "Closed"]

    def _summary_row(label, sub):
        n    = len(sub)
        wins = int(sub["Won"].sum()) if n else 0
        return {
            "Group":      label,
            "Positions":  n,
            "Wins":       wins,
            "Win%":       round(wins / n * 100, 1) if n else 0.0,
            "Avg Entry%": round(sub["Avg Price"].mean() * 100, 1) if n else 0.0,
            "Avg Edge%":  round(sub["Edge"].mean() * 100, 2) if n else 0.0,
            "Total P&L":  round(sub["P&L"].sum(), 2) if n else 0.0,
        }

    summary_df = pd.DataFrame([
        _summary_row("Held to expiry", held),
        _summary_row("Early exits",    closed),
        _summary_row("TOTAL",          df),
    ])

    # By price bucket — held positions only
    bins   = [0, 0.50, 0.65, 0.75, 0.85, 0.92, 1.01]
    labels = ["<50%", "50-65%", "65-75%", "75-85%", "85-92%", ">92%"]
    held_copy = held.copy()
    held_copy["Bucket"] = pd.cut(held_copy["Avg Price"], bins=bins, labels=labels, right=False)
    bk = held_copy.groupby("Bucket", observed=True).agg(
        Bets=("Won", "count"),
        Wins=("Won", "sum"),
        Win_pct=("Won", "mean"),
        Avg_entry=("Avg Price", "mean"),
        PnL=("P&L", "sum"),
    ).reset_index()
    bk["Win_pct"]   = (bk["Win_pct"] * 100).round(1)
    bk["Avg_entry"] = (bk["Avg_entry"] * 100).round(1)
    bk["Edge%"]     = (bk["Win_pct"] - bk["Avg_entry"]).round(1)
    bk["PnL"]       = bk["PnL"].round(2)
    bk.columns      = ["Price Bucket", "Bets", "Wins", "Win%", "Avg Entry%", "P&L", "Edge%"]
    bucket_df = bk

    # By category — all settled
    cat = df.groupby("Category").agg(
        Bets=("Won", "count"),
        Wins=("Won", "sum"),
        Win_pct=("Won", "mean"),
        Avg_entry=("Avg Price", "mean"),
        Avg_edge=("Edge", "mean"),
        PnL=("P&L", "sum"),
    ).reset_index().sort_values("PnL")
    cat["Win_pct"]   = (cat["Win_pct"] * 100).round(1)
    cat["Avg_entry"] = (cat["Avg_entry"] * 100).round(1)
    cat["Avg_edge"]  = (cat["Avg_edge"] * 100).round(2)
    cat["PnL"]       = cat["PnL"].round(2)
    cat.columns      = ["Category", "Bets", "Wins", "Win%", "Avg Entry%", "Avg Edge%", "P&L"]
    category_df = cat.reset_index(drop=True)

    return {"summary": summary_df, "by_bucket": bucket_df, "by_category": category_df}


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

COLORS = {
    "header_blue":  "1F4E79",
    "header_green": "375623",
    "header_teal":  "215868",
    "alt_row":      "EBF3FB",
}


def _style_sheet(ws, header_color: str):
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    alt_fill = PatternFill("solid", fgColor=COLORS["alt_row"])
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float)) else "left")
            if i % 2 == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 35)

    ws.freeze_panes = "A2"


def _add_bar_chart(ws, title, data_col, n_rows, anchor, horizontal=True):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.style = 10
    chart.width = 22
    chart.height = 14
    chart.legend = None

    data = Reference(ws, min_col=data_col, min_row=1, max_row=n_rows + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def _add_line_chart(ws, pivot_df, anchor):
    n_rows = len(pivot_df)
    n_cols = len(pivot_df.columns)

    chart = LineChart()
    chart.title = "Monthly Volume by Category"
    chart.style = 10
    chart.y_axis.title = "Contracts"
    chart.x_axis.title = "Month"
    chart.width = 30
    chart.height = 16

    cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
    for col_idx in range(2, n_cols + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n_rows + 1)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def _style_pnl_sheet(ws, header_color: str):
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    green_fill  = PatternFill("solid", fgColor="C6EFCE")
    red_fill    = PatternFill("solid", fgColor="FFC7CE")
    alt_fill    = PatternFill("solid", fgColor=COLORS["alt_row"])

    # Find P&L column index
    pnl_col = None
    for cell in ws[1]:
        if str(cell.value) == "P&L":
            pnl_col = cell.column
            break

    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float)) else "left")
            if pnl_col and cell.column == pnl_col:
                if isinstance(cell.value, (int, float)):
                    cell.fill = green_fill if cell.value >= 0 else red_fill
                    cell.font = Font(bold=True)
            elif i % 2 == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 50)

    ws.freeze_panes = "A2"


def _write_analysis_sheet(wb, analysis: dict):
    if not analysis:
        return

    from openpyxl.chart import LineChart, Reference as Ref

    ws = wb.create_sheet("Trade Analysis")
    thin = Side(style="thin", color="D0D0D0")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    SEC_COLOR   = "1F4E79"
    HDR_COLOR   = "2E75B6"
    TOTAL_COLOR = "D9E1F2"
    green_fill  = PatternFill("solid", fgColor="C6EFCE")
    red_fill    = PatternFill("solid", fgColor="FFC7CE")
    alt_fill    = PatternFill("solid", fgColor="EBF3FB")
    MONEY_COLS  = {"P&L", "Total P&L"}
    EDGE_COLS   = {"Edge%", "Avg Edge%"}

    def _write_table(df, start_row, pnl_highlight=True):
        """Write header + data rows; return (first_data_row, last_data_row)."""
        n_cols = len(df.columns)
        # Header row
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(start_row, j, col)
            c.font      = Font(bold=True, color="FFFFFF", size=10)
            c.fill      = PatternFill("solid", fgColor=HDR_COLOR)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = bdr
        ws.row_dimensions[start_row].height = 20
        first_data = start_row + 1
        cur = first_data
        n_rows = len(df)
        for i, (_, rdata) in enumerate(df.iterrows()):
            is_total = str(rdata.iloc[0]) == "TOTAL"
            for j, val in enumerate(rdata, 1):
                col_name = df.columns[j - 1]
                c = ws.cell(cur, j, val)
                c.border    = bdr
                c.alignment = Alignment(
                    horizontal="right" if isinstance(val, (int, float)) else "left"
                )
                if is_total:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor=TOTAL_COLOR)
                elif pnl_highlight and col_name in MONEY_COLS and isinstance(val, (int, float)):
                    c.fill = green_fill if val >= 0 else red_fill
                    c.font = Font(bold=True)
                elif pnl_highlight and col_name in EDGE_COLS and isinstance(val, (int, float)):
                    c.fill = green_fill if val > 0 else (red_fill if val < 0 else PatternFill())
                elif i % 2 == 1:
                    c.fill = alt_fill
            cur += 1
        return first_data, cur - 1

    def _section_header(title, row, n_cols):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        c = ws.cell(row, 1, title)
        c.font      = Font(bold=True, color="FFFFFF", size=11)
        c.fill      = PatternFill("solid", fgColor=SEC_COLOR)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 22

    row = 1
    summary = analysis["summary"]
    bucket  = analysis["by_bucket"]
    categ   = analysis["by_category"]

    # Section 1: Held vs Early Exit
    _section_header("Held to Expiry vs Early Exits", row, len(summary.columns))
    row += 1
    _, _ = _write_table(summary, row)
    row += len(summary) + 3

    # Section 2: By Price Bucket
    bucket_sec_row = row
    _section_header("Win Rate vs Implied Probability — Held to Expiry", row, len(bucket.columns))
    row += 1
    bk_hdr_row = row
    bk_first, bk_last = _write_table(bucket, row)
    row = bk_last + 3

    # Section 3: By Category
    _section_header("Performance by Category", row, len(categ.columns))
    row += 1
    _write_table(categ, row)

    # Auto-size columns
    for col in ws.columns:
        col_cells = [c for c in col if c.value is not None]
        if not col_cells:
            continue
        max_len = max(len(str(c.value)) for c in col_cells)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 28)

    ws.freeze_panes = "A2"

    # Calibration chart: Win% (col 4) vs Avg Entry% (col 5) by price bucket
    if not bucket.empty:
        n = len(bucket)
        chart = LineChart()
        chart.title  = "Calibration: Actual Win% vs Implied Probability"
        chart.style  = 10
        chart.y_axis.title = "Percent"
        chart.x_axis.title = "Entry Price Bucket"
        chart.width  = 24
        chart.height = 14

        cats       = Ref(ws, min_col=1, min_row=bk_first, max_row=bk_last)
        win_data   = Ref(ws, min_col=4, min_row=bk_hdr_row, max_row=bk_last)  # Win% with header
        entry_data = Ref(ws, min_col=5, min_row=bk_hdr_row, max_row=bk_last)  # Avg Entry% with header

        chart.add_data(win_data,   titles_from_data=True)
        chart.add_data(entry_data, titles_from_data=True)
        chart.set_categories(cats)

        chart.series[0].graphicalProperties.line.solidFill = "00B050"  # green = actual win%
        chart.series[0].graphicalProperties.line.width = 25000
        chart.series[1].graphicalProperties.line.solidFill = "4472C4"  # blue = implied
        chart.series[1].graphicalProperties.line.width = 20000

        anchor_col = len(bucket.columns) + 2
        ws.add_chart(chart, f"{get_column_letter(anchor_col)}{bucket_sec_row}")


def export_excel(cat_df, time_df, active_df, settled_df=None, open_df=None, analysis=None):
    sheets = {"By Category": cat_df, "Over Time": time_df, "Active Now": active_df}
    if settled_df is not None and not settled_df.empty:
        sheets["Settled P&L"] = settled_df
    if open_df is not None and not open_df.empty:
        sheets["Open Positions"] = open_df

    with pd.ExcelWriter(EXCEL_OUT, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(EXCEL_OUT)

    ws_cat = wb["By Category"]
    _style_sheet(ws_cat, COLORS["header_blue"])
    _add_bar_chart(ws_cat, "Total Volume by Category", 3, len(cat_df), "E2", horizontal=True)

    ws_time = wb["Over Time"]
    _style_sheet(ws_time, COLORS["header_green"])
    _add_line_chart(ws_time, time_df, "A" + str(len(time_df) + 4))

    ws_active = wb["Active Now"]
    _style_sheet(ws_active, COLORS["header_teal"])
    _add_bar_chart(ws_active, "24h Volume by Category", 3, len(active_df), "F2", horizontal=True)

    if "Settled P&L" in wb.sheetnames:
        _style_pnl_sheet(wb["Settled P&L"], COLORS["header_blue"])
        # Add totals row
        ws = wb["Settled P&L"]
        last_row = ws.max_row + 2
        ws.cell(last_row, 1, "TOTAL").font = Font(bold=True)
        pnl_col = next((c.column for c in ws[1] if c.value == "P&L"), None)
        if pnl_col:
            col_letter = get_column_letter(pnl_col)
            ws.cell(last_row, pnl_col).value = f"=SUM({col_letter}2:{col_letter}{ws.max_row - 1})"
            ws.cell(last_row, pnl_col).font = Font(bold=True)

    if "Open Positions" in wb.sheetnames:
        _style_pnl_sheet(wb["Open Positions"], COLORS["header_teal"])

    if analysis:
        _write_analysis_sheet(wb, analysis)

    wb.save(EXCEL_OUT)
    print(f"\nDashboard saved: {EXCEL_OUT}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    series_map = build_series_category_map()
    with KalshiClient.from_env() as client:
        settled_rows = pull_settled(client, series_map)
        open_rows    = pull_open_24h(client, series_map)
        settled_pnl, open_positions = build_trade_pnl(client)

    if not settled_rows and not open_rows and settled_pnl.empty:
        print("No data found.")
        return

    cat_df    = category_totals(settled_rows) if settled_rows else pd.DataFrame()
    time_df   = volume_over_time(settled_rows) if settled_rows else pd.DataFrame()
    active_df = active_now(open_rows) if open_rows else pd.DataFrame()
    analysis  = build_trade_analysis(settled_pnl)

    print("\n--- By Category (settled) ---")
    print(cat_df.to_string(index=False))
    print("\n--- Active Now (24h volume) ---")
    print(active_df.to_string(index=False))

    if not settled_pnl.empty:
        total_pnl = settled_pnl["P&L"].sum()
        wins = (settled_pnl["P&L"] > 0).sum()
        print(f"\n--- Trade P&L ({len(settled_pnl)} settled, {wins}W/{len(settled_pnl)-wins}L, total {total_pnl:+.2f}) ---")
        print(settled_pnl[["Ticker","Result","Net YES","P&L"]].to_string(index=False))
    if analysis:
        print("\n--- Trade Analysis ---")
        print(analysis["summary"].to_string(index=False))
        print()
        print(analysis["by_bucket"].to_string(index=False))
    if not open_positions.empty:
        print(f"\n--- Open Positions ({len(open_positions)}) ---")
        print(open_positions[["Ticker","Net YES","Cost Basis"]].to_string(index=False))

    export_excel(cat_df, time_df, active_df, settled_pnl, open_positions, analysis)


if __name__ == "__main__":
    main()
